"""Automated Integration & End-to-End Tests for CareerLens AI."""

import io
from fastapi.testclient import TestClient
from backend.main import app
from backend.sample_data import SAMPLE_RESUMES, SAMPLE_JDS
import docx
import pypdf

client = TestClient(app)

def test_health():
    response = client.get("/api/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "healthy"
    print("PASS: /api/health")

def test_samples():
    response = client.get("/api/samples")
    assert response.status_code == 200
    data = response.json()
    assert "sde_fresher" in data["resumes"]
    assert "google_sde" in data["job_descriptions"]
    print("PASS: /api/samples")

def test_analyze_with_text():
    resume_text = SAMPLE_RESUMES["sde_fresher"]["text"]
    jd_text = SAMPLE_JDS["google_sde"]["text"]
    
    response = client.post(
        "/api/analyze",
        data={
            "resume_text": resume_text,
            "jd_text": jd_text,
            "target_role": "Full-Stack SDE-1"
        }
    )
    assert response.status_code == 200
    data = response.json()
    assert "match_analysis" in data
    assert "career_critique" in data
    assert "roadmap" in data
    assert "interview_questions" in data
    
    match = data["match_analysis"]
    print(f"Overall Match Score: {match['overall_score']}%")
    print(f"Matching Skills ({len(match['matching_skills'])}): {[s['name'] for s in match['matching_skills']]}")
    print(f"Missing Skills ({len(match['missing_skills'])}): {[s['name'] for s in match['missing_skills']]}")
    assert len(match["matching_skills"]) > 0
    assert len(match["missing_skills"]) > 0
    assert len(data["interview_questions"]) >= 10
    assert len(data["roadmap"]) == 4
    print("PASS: /api/analyze (Text Mode)")

def test_mock_interview_flow():
    # 1. Start mock session
    start_resp = client.post(
        "/api/mock-interview/start",
        json={
            "role": "SDE-1",
            "candidate_name": "Aarav Sharma",
            "questions": [
                {
                    "id": 1,
                    "category": "Technical",
                    "question": "How do you handle database transaction boundaries in Java?",
                    "why_asked": "Assesses backend integrity",
                    "expected_keywords": ["ACID", "Isolation Levels", "@Transactional"],
                    "model_answer_tip": "Mention ACID and @Transactional"
                }
            ]
        }
    )
    assert start_resp.status_code == 200
    session_data = start_resp.json()
    session_id = session_data["session_id"]
    assert session_id is not None
    print(f"Mock Interview Session Initialized: {session_id}")

    # 2. Submit candidate answer
    submit_resp = client.post(
        "/api/mock-interview/submit",
        json={
            "session_id": session_id,
            "candidate_answer": "In Spring Boot and Java, I use the @Transactional annotation on service methods to define transaction boundaries. This ensures ACID compliance and automatically rolls back if an unhandled RuntimeException occurs."
        }
    )
    assert submit_resp.status_code == 200
    eval_data = submit_resp.json()
    assert "evaluated_entry" in eval_data
    evaluation = eval_data["evaluated_entry"]["evaluation"]
    print(f"Answer Score: {evaluation['score']}/10, Verdict: {evaluation['verdict']}")
    assert evaluation["score"] >= 6
    assert len(evaluation["strengths"]) > 0
    assert len(evaluation["exemplary_answer"]) > 0

    # 3. Fetch report card
    report_resp = client.get(f"/api/mock-interview/report/{session_id}")
    assert report_resp.status_code == 200
    report = report_resp.json()
    assert report["completed_questions"] == 1
    assert report["average_score"] > 0
    print("PASS: Mock Interview Full Flow")

def test_bullet_rewrite():
    resp = client.post(
        "/api/rewrite-bullet",
        json={
            "bullet_point": "Built an API for handling orders using Spring Boot and SQL.",
            "target_keywords": ["REST API", "PostgreSQL", "Scalability"]
        }
    )
    assert resp.status_code == 200
    data = resp.json()
    assert len(data["rewrites"]) == 3
    print("PASS: /api/rewrite-bullet")

def test_file_upload_parsing():
    # Create an in-memory docx file
    doc = docx.Document()
    doc.add_heading("TEST CANDIDATE", 0)
    doc.add_paragraph("Skills: Python, Django, PostgreSQL, Docker, AWS")
    doc.add_paragraph("Experience: Built cloud microservices.")
    docx_stream = io.BytesIO()
    doc.save(docx_stream)
    docx_stream.seek(0)

    jd_text = "Looking for a Python Backend Developer skilled in Python, Django, PostgreSQL, Docker, and AWS."
    
    resp = client.post(
        "/api/analyze",
        data={"jd_text": jd_text, "target_role": "Python Developer"},
        files={"resume_file": ("test_resume.docx", docx_stream.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert resp.status_code == 200
    res_data = resp.json()
    assert res_data["match_analysis"]["overall_score"] >= 70
    print("PASS: File Upload Parsing (DOCX)")

if __name__ == "__main__":
    test_health()
    test_samples()
    test_analyze_with_text()
    test_mock_interview_flow()
    test_bullet_rewrite()
    test_file_upload_parsing()
    print("\nALL 6 ENDPOINT & INTEGRATION TESTS PASSED PERFECTLY! [SUCCESS]")
